import json,io,csv,os
from datetime import datetime,date,time,timedelta
from zoneinfo import ZoneInfo
from fastapi import APIRouter,Depends,HTTPException,Request
from fastapi.responses import StreamingResponse
from pydantic import Field
from . import store as s,security as sec,tasks
from .models import Strict,Record
from .main import db_dep,user_dep,idem

router=APIRouter()
class ReportRequest(Strict):
    start:str
    end:str
    title:str=Field(default='Department work report',max_length=200)
def factual(db,start,end,title):
    settings=json.loads(db.execute('SELECT body FROM settings').fetchone()[0]);tz=ZoneInfo(settings['timezone'])
    try:
        lo=datetime.combine(date.fromisoformat(start),time.min,tz)
        hi=datetime.combine(date.fromisoformat(end)+timedelta(days=1),time.min,tz)
    except ValueError: raise HTTPException(422,'Use valid start/end dates.')
    if hi<=lo: raise HTTPException(422,'End precedes start.')
    sections={k:[] for k in ['Completed work','Awaiting manager review','Ongoing work','Blocked / delayed','Routine checks','Priority attention','Next-period plan','Member contributions','Decisions / resources required']}
    rows=[];sources=[]
    for task in s.listing(db,'tasks')+s.listing(db,'tasks',True):
        events=[dict(r) for r in db.execute('SELECT * FROM history WHERE entity=? ORDER BY seq',(task['id'],))]
        before=[e for e in events if datetime.fromisoformat(e['at'])<hi]
        retrospective=False
        if not before:
            if task.get('occurrence') and start<=task['occurrence']<=end:
                occurrence_event=next((e for e in events if e['action']=='schedule occurrence'),None)
                if not occurrence_event:continue
                snapshot=json.loads(occurrence_event['snapshot']);retrospective=True
            else:continue
        else:snapshot=json.loads(before[-1]['snapshot'])
        changed=any(lo<=datetime.fromisoformat(e['at'])<hi for e in before)
        due=datetime.fromisoformat(snapshot['due']) if snapshot['due'] else None
        active=snapshot['status'] not in ['Completed','Cancelled']
        if not changed and not active and not (due and lo<=due<hi): continue
        approval='Manager approved' if snapshot.get('approved_at') else 'Member reported; not manager approved' if snapshot['status']=='Completed' else snapshot['status']
        row={'no':len(rows)+1,'work_scope':snapshot['title'],'activity':snapshot.get('completion_summary') or snapshot.get('latest_note') or snapshot['description'],'status':approval,'task_id':task['id'],'task_number':task['number'],'owner':snapshot['owner'],'source_version':snapshot['version'],'due':snapshot['due'],'not_applicable':snapshot.get('not_applicable',False)}
        if row['not_applicable']: row['status']='Not applicable — '+approval
        if retrospective:row['status']='Missed occurrence — created after period end'
        rows.append(row);sources.append(snapshot)
        key='Completed work' if snapshot['status']=='Completed' else 'Awaiting manager review' if snapshot['status']=='Submitted for Review' else 'Blocked / delayed' if snapshot['status']=='Waiting' else 'Ongoing work'
        if snapshot['status']!='Cancelled': sections[key].append(task['id'])
        if snapshot.get('schedule_id'): sections['Routine checks'].append(task['id'])
        if active and snapshot['priority'] in ['High','Critical']: sections['Priority attention'].append(task['id'])
        if active: sections['Next-period plan'].append(task['id'])
        if snapshot.get('blocker'): sections['Decisions / resources required'].append(task['id'])
    contributors={}
    ids={r['task_id'] for r in rows}
    for event in db.execute('SELECT * FROM history WHERE actor IS NOT NULL'):
        if event['entity'] in ids and lo<=datetime.fromisoformat(event['at'])<hi:
            contributors.setdefault(str(event['actor']),set()).add(event['entity'])
    sections['Member contributions']={k:sorted(v) for k,v in contributors.items()}
    return {'title':title,'start':start,'end':end,'timezone':settings['timezone'],'rows':rows,'sections':sections,'sources':sources,'summary':'','decisions':'','finalized':False,'inclusion_rules':'Inclusive business dates. One row per task: activity within the period, due within the period, or still active at period end. State is reconstructed from history at period end. Later changes are excluded. Backfilled scheduled occurrences due in the period are labelled missed and created after period end; later completion is not counted. Source snapshots and versions are retained. Contribution counts identify involvement, not employee performance.'}
@router.get('/api/reports')
def reports(user=Depends(user_dep),db=Depends(db_dep)):
    sec.manager(user);return s.listing(db,'reports')
@router.post('/api/reports')
def generate(body:ReportRequest,request:Request,user=Depends(user_dep),db=Depends(db_dep)):
    sec.manager(user)
    return idem(db,user,request,body.model_dump(),lambda:s.insert(db,'reports',factual(db,body.start,body.end,body.title),user['id']))
@router.put('/api/reports/{id}')
def edit_report(id:str,body:Record,request:Request,user=Depends(user_dep),db=Depends(db_dep)):
    sec.manager(user)
    def write():
        doc=s.get(db,id)
        if not doc or doc['kind']!='reports': raise HTTPException(404,'Report not found.')
        tasks.version(doc,body.version)
        if doc['finalized']: raise HTTPException(409,'Finalized snapshots cannot be edited. Generate a new report.')
        allowed={'title','summary','decisions','rows','finalized'}
        if set(body.data)-allowed: raise HTTPException(422,'Unsupported report field.')
        if 'rows' in body.data:
            rows=body.data['rows']
            if not isinstance(rows,list) or len(rows)!=len(doc['rows']): raise HTTPException(422,'Preserve the source rows.')
            for old,new in zip(doc['rows'],rows):
                if set(old)!=set(new) or any(old[k]!=new[k] for k in old if k not in ['work_scope','activity','status']): raise HTTPException(422,'Source references cannot be changed.')
                if not all(isinstance(new[k],str) and len(new[k])<=20000 for k in ['work_scope','activity','status']): raise HTTPException(422,'Invalid report text.')
        for k in ['title','summary','decisions']:
            if k in body.data and (not isinstance(body.data[k],str) or len(body.data[k])>50000): raise HTTPException(422,'Invalid report text.')
        if 'finalized' in body.data and not isinstance(body.data['finalized'],bool): raise HTTPException(422,'Invalid finalization value.')
        doc.update(body.data)
        if doc['finalized']: doc['finalized_at']=s.now();doc['finalized_by']=user['id']
        return s.save(db,doc,user['id'],'finalize report' if doc['finalized'] else 'report revision',body.reason)
    return idem(db,user,request,body.model_dump(),write)
@router.get('/api/reports/{id}/excel')
def excel(id:str,user=Depends(user_dep),db=Depends(db_dep)):
    sec.manager(user);doc=s.get(db,id)
    if not doc or doc['kind']!='reports': raise HTTPException(404,'Report not found.')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();ws=book.active;ws.title='Work report'
    def safe(v): return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
    ws.append([safe(doc['title'])]);ws.append([doc['start'],doc['end'],doc['timezone'],'Finalized' if doc['finalized'] else 'Draft']);ws.append([safe(doc['summary'])]);ws.append(['No.','Work Scope','Activity','Status','Source task','Source version'])
    for row in doc['rows']: ws.append([row['no'],safe(row['work_scope']),safe(row['activity']),safe(row['status']),row['task_number'],row['source_version']])
    ws.freeze_panes='A5';ws.auto_filter.ref=f'A4:F{max(ws.max_row,4)}'
    for c in ws[4]: c.fill=PatternFill('solid',fgColor='173F43');c.font=Font(color='FFFFFF',bold=True)
    for i,width in enumerate([7,42,64,38,18,16],1): ws.column_dimensions[get_column_letter(i)].width=width
    for row in ws.iter_rows(min_row=5):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=60
    source=book.create_sheet('Sources');source.append(['Task','Source link','Version','Owner','Status','Approved at'])
    for task in doc['sources']:
        link=os.getenv('PUBLIC_ORIGIN','http://127.0.0.1:8000').rstrip('/')+f"/#task/{task['id']}"
        source.append([task['number'],link,task['version'],task['owner'],task['status'],task.get('approved_at') or ''])
        source.cell(source.max_row,2).hyperlink=link
        source.cell(source.max_row,2).style='Hyperlink'
    source.append(['Inclusion rules',doc['inclusion_rules']]);source.column_dimensions['A'].width=24;source.column_dimensions['B'].width=90
    stream=io.BytesIO();book.save(stream);stream.seek(0)
    return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="kelun-report-{doc["start"]}.xlsx"'})
@router.get('/api/export/tasks')
def export_tasks(user=Depends(user_dep),db=Depends(db_dep)):
    sec.manager(user);buf=io.StringIO();writer=csv.writer(buf);fields=['number','title','owner','status','priority','due','created','updated','completion_summary','test_results','approved_at','archived']
    writer.writerow(fields)
    for task in s.listing(db,'tasks')+s.listing(db,'tasks',True): writer.writerow([("'"+str(task.get(k,''))) if str(task.get(k,'')).startswith(('=','+','-','@')) else task.get(k,'') for k in fields])
    return StreamingResponse(iter([buf.getvalue().encode('utf-8-sig')]),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="kelun-tasks.csv"'})
