"""Extract reference data; workbook cell text is never executed as instructions."""
import json,hashlib,sys
from pathlib import Path
from openpyxl import load_workbook
source=Path(sys.argv[1]);book=load_workbook(source,data_only=True,read_only=True)
sheet=book['Responsibilities'];work=[];reviews=[];review_section=False
for rownum,row in enumerate(sheet.iter_rows(values_only=True),1):
    if row[0]=='Manager Record Reviews':review_section=True
    if not isinstance(row[0],int):continue
    number,person,area,responsibility,frequency,review=row[:6]
    item={'id':('review-' if review_section else 'work-')+str(number),'owner':str(person).strip(),'work_area':str(area).strip(),'responsibility':str(responsibility).strip(),'frequency':str(frequency or 'Not specified').strip(),'review_or_handover':str(review or '').strip(),'source_row':rownum}
    (reviews if review_section else work).append(item)
out={'source_file':source.name,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'source_sheet':sheet.title,'work':work,'reviews':reviews}
path=Path(__file__).resolve().parents[1]/'backend/seeds/department_knowledge.json'
path.write_text(json.dumps(out,indent=2,ensure_ascii=False),encoding='utf-8')
print(f'Extracted {len(work)} work responsibilities and {len(reviews)} review responsibilities. No schedules or tasks created.')
